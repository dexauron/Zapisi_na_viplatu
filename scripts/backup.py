"""Ночная резервная копия: все таблицы Supabase → Excel (лист на таблицу) + JSON.
Запуск: SUPABASE_URL=... SUPABASE_SECRET_KEY=... python scripts/backup.py <папка>
Ключ — секретный (sb_secret_…), он обходит RLS; хранится только в секретах GitHub."""
import json
import os
import sys
import urllib.request
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

URL = os.environ["SUPABASE_URL"].rstrip("/")
KEY = os.environ["SUPABASE_SECRET_KEY"]
OUT = sys.argv[1] if len(sys.argv) > 1 else "backup"

# таблица → (лист, порядок сортировки)
TABLES = {
    "orders": ("Заказы", "order_date.desc"),
    "payouts": ("Выплаты", "payout_date.desc"),
    "reps": ("Представители", "full_name"),
    "stores": ("Магазины", "name"),
    "staff": ("Сотрудники", "full_name"),
    "rep_access": ("Ссылки ТП", "rep_id"),
    "audit_log": ("История", "at.desc"),
    "admins": ("Владельцы", "email"),
}
PAGE = 1000


def fetch(table, order):
    rows, start = [], 0
    while True:
        req = urllib.request.Request(
            f"{URL}/rest/v1/{table}?select=*&order={order}",
            headers={"apikey": KEY, "Authorization": f"Bearer {KEY}", "Range-Unit": "items",
                     "Range": f"{start}-{start + PAGE - 1}"},
        )
        with urllib.request.urlopen(req, timeout=60) as r:
            chunk = json.load(r)
        rows += chunk
        if len(chunk) < PAGE:
            return rows
        start += PAGE


def main():
    os.makedirs(OUT, exist_ok=True)
    data = {t: fetch(t, o) for t, (_, o) in TABLES.items()}
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    with open(f"{OUT}/sverka-{stamp}.json", "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1, default=str)

    reps = {r["id"]: r["full_name"] for r in data["reps"]}
    wb = Workbook()
    wb.remove(wb.active)
    for t, (title, _) in TABLES.items():
        ws = wb.create_sheet(title)
        rows = data[t]
        cols = list(rows[0].keys()) if rows else ["(пусто)"]
        if "rep_id" in cols:
            cols.insert(cols.index("rep_id") + 1, "rep_name")
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            r = dict(r, rep_name=reps.get(r.get("rep_id"), ""))
            ws.append([json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for v in (r.get(c) for c in cols)])
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(40, max(10, len(c) + 4))
        ws.freeze_panes = "A2"
    wb.save(f"{OUT}/sverka-{stamp}.xlsx")
    print("\n".join(f"{TABLES[t][0]}: {len(v)}" for t, v in data.items()))


if __name__ == "__main__":
    main()
